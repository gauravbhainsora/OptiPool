import csv
import io
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .pods import selected_pods


def iso(value):
    return datetime.fromtimestamp(value, UTC).isoformat() if value else "Not fetched"


def pool_rows(inventory, snapshots=None):
    rows = []
    for pool in inventory:
        saved = (snapshots or {}).get(pool["id"])
        summary = pool.get("snapshot_summary")
        if saved:
            from .store import overview

            summary = overview(saved)["summary"]
        pool = {
            **pool,
            **(saved["pool"] if saved else pool.get("snapshot_pool", {})),
            "available": pool["available"],
        }
        rows.append(
            {
                "Project": pool["project"],
                "Cluster": pool["cluster"],
                "Node pool": pool["name"],
                "Environment": pool["environment"],
                "Location": pool["location"],
                "Machine type": pool["machine_type"],
                "Available": pool["available"],
                "Nodes": summary["node_count"] if summary else "Unknown",
                "Autoscaler min": pool["autoscaler_min"],
                "Autoscaler max": pool["autoscaler_max"],
                "Allocatable CPU cores": summary["cpu_allocatable"] if summary else None,
                "Requested CPU cores": summary["cpu"] if summary else None,
                "Allocatable memory GiB": summary["memory_allocatable"] / 1024**3 if summary else None,
                "Requested memory GiB": summary["memory"] / 1024**3 if summary else None,
                "Policy violations": summary["violations"] if summary else None,
                "Snapshot UTC": iso(saved["saved_at"] if saved else pool.get("snapshot_at")),
                "Snapshot ID": saved["id"] if saved else pool.get("snapshot_id") or "",
                "Synthetic": pool.get("synthetic", False),
            }
        )
    return rows


def workload_rows(snapshot, usage):
    rows = []
    for w in snapshot["workloads"]:
        metric = usage["workloads"].get(w["id"], {}) if w["kind"] == "Deployment" else {}
        rows.append(
            {
                "Project": snapshot["pool"]["project"],
                "Cluster": snapshot["pool"]["cluster"],
                "Node pool": snapshot["pool"]["name"],
                "Namespace": w["namespace"],
                "Kind": w["kind"],
                "Workload": w["name"],
                "Services": ", ".join(w["services"]),
                "Desired replicas (cluster)": w["desired"],
                "Running replicas (cluster)": w["running"],
                "Pending replicas (cluster)": w["pending"],
                "Replicas on pool": w["pool_replicas"],
                "CPU requests on pool": w["requested"]["cpu"],
                "Memory requests on pool GiB": w["requested"]["memory"] / 1024**3,
                "Highest replica p99 CPU cores": metric.get("cpu_p99"),
                "Highest replica p99 memory GiB": metric["memory_p99"] / 1024**3
                if metric.get("memory_p99") is not None
                else None,
                "Historical replicas observed": metric.get("historical_replicas_observed", 0),
                "Historical replica UIDs": ", ".join(metric.get("historical_replica_uids", [])),
                "Historical attribution gaps": "; ".join(usage.get("attribution_gaps", [])),
                "CPU p99 replica": metric.get("cpu_replica"),
                "Memory p99 replica": metric.get("memory_replica"),
                "Window days": usage["days"],
                "Metrics cutoff UTC": iso(usage["cutoff"]),
                "Metric source": usage["source"],
                "Memory definition": usage["memory_definition"],
                "Coverage": usage["coverage"],
                "Snapshot UTC": iso(snapshot["saved_at"]),
                "Snapshot ID": snapshot["id"],
                "Synthetic": snapshot["pool"].get("synthetic", False),
            }
        )
    return rows


def pod_rows(snapshot, search="", usage=None):
    rows = []
    for p in selected_pods(snapshot):
        metric = (usage or {}).get("pods", {}).get(p["uid"], {})
        if (
            search
            and search.lower() not in f"{p['name']} {p['namespace']} {p['workload']} {p['node']}".lower()
        ):
            continue
        rows.append(
            {
                "Project": snapshot["pool"]["project"],
                "Cluster": snapshot["pool"]["cluster"],
                "Node pool": snapshot["pool"]["name"],
                "Pod": p["name"],
                "Pod UID": p["uid"],
                "Namespace": p["namespace"],
                "Workload": p["workload"],
                "Node": p["node"],
                "Phase": p["phase"],
                "Ready": p["ready"],
                "Restarts": p["restarts"],
                "CPU request cores": p["requested"]["cpu"],
                "Memory request GiB": p["requested"]["memory"] / 1024**3,
                "CPU limit cores": p["limits"]["cpu"],
                "Memory limit GiB": p["limits"]["memory"] / 1024**3,
                "Window days": (usage or {}).get("days"),
                "CPU latest cores": metric.get("cpu_latest"),
                "Memory latest GiB": metric["memory_latest"] / 1024**3
                if metric.get("memory_latest") is not None
                else None,
                "CPU p99 cores": metric.get("cpu_p99"),
                "Memory p99 GiB": metric["memory_p99"] / 1024**3
                if metric.get("memory_p99") is not None
                else None,
                "CPU sample UTC": iso(metric.get("cpu_sample_at")),
                "Memory sample UTC": iso(metric.get("memory_sample_at")),
                "CPU sample coverage": metric.get("cpu_coverage"),
                "Memory sample coverage": metric.get("memory_coverage"),
                "OOMKilled observed": p["oom_killed"] if p["status_available"] else None,
                "Status captured": p["status_available"],
                "Created UTC": iso(p["created"]),
                "Snapshot UTC": iso(snapshot["saved_at"]),
                "Snapshot ID": snapshot["id"],
                "Synthetic": snapshot["pool"].get("synthetic", False),
            }
        )
    return rows


def safe_cell(value):
    # Prevent spreadsheet formula injection from untrusted cloud labels/names.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def render(rows, format, table):
    columns = list(rows[0]) if rows else ["No saved rows"]
    if format == "csv":
        out = io.StringIO(newline="")
        writer = csv.writer(out)
        writer.writerow(columns)
        writer.writerows([[safe_cell(row.get(c)) for c in columns] for row in rows])
        return out.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = table
    sheet.append(columns)
    for row in rows:
        sheet.append([safe_cell(row.get(c)) for c in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B37")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = min(55, max(14, len(str(col[0].value)) + 2))
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
