import csv
import io
import threading
import time
from copy import deepcopy

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from optipool.app import create_app
from optipool.demo import DemoCollector
from optipool.exports import render
from optipool.store import Store


class CountingCollector(DemoCollector):
    def __init__(self, settings):
        super().__init__(settings)
        self.calls = 0
        self.fail = False
        self.block = None

    def collect(self, pool):
        self.calls += 1
        if self.block:
            self.block.wait(3)
        if self.fail:
            raise ValueError("Core collection failed")
        return super().collect(pool)


def wait_job(client, job_id):
    for _ in range(300):
        r = client.get(f"/api/jobs/{job_id}").json()
        if r["state"] not in {"running", "queued"}:
            return r
        time.sleep(0.01)
    pytest.fail("job did not complete")


def initialized(client):
    for _ in range(300):
        inv = client.get("/api/inventory").json()
        if inv["initialized"]:
            return inv
        time.sleep(0.01)
    pytest.fail("initial import did not complete")


def test_full_flow_saved_reads_exports_and_scoped_failed_refresh(settings):
    collector = CountingCollector(settings)
    with TestClient(create_app(settings, collector)) as client:
        inv = initialized(client)
        assert len(inv["pools"]) == 8 and collector.calls == 8
        a, b = inv["pools"][:2]
        snap_a = client.get("/api/snapshot", params={"pool_id": a["id"]}).json()
        snap_b = client.get("/api/snapshot", params={"pool_id": b["id"]}).json()
        for days in [1, 5, 10]:
            r = client.get("/api/snapshot", params={"pool_id": a["id"], "days": days})
            assert r.status_code == 200 and r.json()["usage"]["days"] == days
        assert len(snap_a["pods"]) == 10
        for table in ["node-pools", "workloads", "pods"]:
            for fmt in ["csv", "xlsx"]:
                r = client.get(f"/api/export/{table}", params={"pool_id": a["id"], "format": fmt, "days": 5})
                assert r.status_code == 200
                if fmt == "xlsx":
                    book = load_workbook(io.BytesIO(r.content))
                    assert book.active.max_row > 1
                    assert "Snapshot ID" in [c.value for c in book.active[1]]
                else:
                    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
                    assert rows and "Snapshot ID" in rows[0]
                    if table == "workloads":
                        assert rows[0]["Window days"] == "5"
        r = client.post(
            "/api/estimate", params={"pool_id": a["id"]}, json={"cpu": "1", "memory": "1Gi", "replicas": 10}
        )
        assert r.status_code == 200 and r.json()["snapshot_id"] == snap_a["id"]
        assert r.json()["target_pool"]["id"] == a["id"]
        assert collector.calls == 8  # Views, exports and estimates never call the collector.
        collector.block = threading.Event()
        job = client.post("/api/refresh/pool", params={"pool_id": a["id"]}, json={}).json()
        same = client.post("/api/refresh/pool", params={"pool_id": a["id"]}, json={}).json()
        assert job["id"] == same["id"]
        collector.block.set()
        assert wait_job(client, job["id"])["state"] == "succeeded"
        assert client.get("/api/snapshot", params={"pool_id": b["id"]}).json()["id"] == snap_b["id"]
        latest = client.get("/api/snapshot", params={"pool_id": a["id"]}).json()["id"]
        assert latest != snap_a["id"]
        collector.fail = True
        job = client.post("/api/refresh/pool", params={"pool_id": a["id"]}, json={}).json()
        assert wait_job(client, job["id"])["state"] == "failed"
        assert client.get("/api/snapshot", params={"pool_id": a["id"]}).json()["id"] == latest
        previous_calls = collector.calls
        job = client.post("/api/refresh/inventory", json={}).json()
        wait_job(client, job["id"])
        assert collector.calls == previous_calls
    collector.fail = False
    with TestClient(create_app(settings, collector)) as client:
        assert initialized(client)["pools"][0]["snapshot_id"] == latest
        assert collector.calls == previous_calls


def test_store_backup_and_atomic_publication(settings, saved, tmp_path):
    store = Store(settings.database)
    pool = saved["pool"]
    store.publish_inventory([pool], [pool["project"]], [])
    sid = store.publish_snapshot(pool["id"], saved)
    broken = deepcopy(saved)
    broken["core_complete"] = False
    with pytest.raises(ValueError):
        store.publish_snapshot(pool["id"], broken)
    assert store.snapshot(pool_id=pool["id"])["id"] == sid
    backup = str(tmp_path / "backup.db")
    store.backup(backup)
    restored = Store(backup)
    assert restored.snapshot(pool_id=pool["id"])["id"] == sid
    restored.close()
    store.close()


def test_discovery_does_not_delete_failed_project(settings, saved):
    s = Store(settings.database)
    pool = saved["pool"]
    s.publish_inventory([pool], [pool["project"]], [])
    s.publish_inventory([], [], [{"scope": pool["project"], "error": "forbidden"}])
    assert s.pool(pool["id"])["available"]
    s.publish_inventory([], [pool["project"]], [])
    assert not s.pool(pool["id"])["available"]
    s.close()


def test_export_formula_injection():
    rows = [{"Name": '=HYPERLINK("bad")', "Value": 2}]
    content, _ = render(rows, "xlsx", "table")
    cell = load_workbook(io.BytesIO(content)).active["A2"]
    assert cell.data_type == "s" and cell.value.startswith("'")
    content, _ = render(rows, "csv", "table")
    assert "'=HYPERLINK" in content.decode("utf-8-sig")


def test_history_pins_snapshot_retention(settings, saved):
    s = Store(settings.database)
    first = s.publish_snapshot(saved["pool"]["id"], saved)
    second = s.publish_snapshot(saved["pool"]["id"], saved)
    s.save_estimate(first, {}, {}, "test")
    with s.db:
        s.db.execute("UPDATE snapshots SET created=1")
    s.prune()
    assert s.snapshot(snapshot_id=first)
    assert s.snapshot(snapshot_id=second)
    s.close()


def test_single_owner_and_future_schema_rejected(settings):
    s = Store(settings.database)
    with pytest.raises(ValueError, match="already has"):
        Store(settings.database)
    with s.db:
        s.db.execute("UPDATE schema_version SET version=2")
    s.close()
    with pytest.raises(ValueError, match="schema version"):
        Store(settings.database)


def test_json_boundary_and_invalid_manifest(settings):
    with TestClient(create_app(settings)) as client:
        pool = initialized(client)["pools"][0]
        assert client.post("/api/refresh/inventory", data={"x": "y"}).status_code == 415
        assert (
            client.post(
                "/api/refresh/inventory", json={}, headers={"Origin": "https://foreign.invalid"}
            ).status_code
            == 403
        )
        assert (
            client.post(
                "/api/estimate", params={"pool_id": pool["id"]}, json={"manifest": "invalid: ["}
            ).status_code
            == 422
        )


def test_saved_source_metadata_reaches_authenticated_api_and_exports_without_reads(settings, monkeypatch):
    import json

    import jwt
    from cryptography.hazmat.primitives.asymmetric import ec
    from google.oauth2 import id_token

    settings.auth_mode = "iap"
    settings.iap_audience = "/projects/123/global/backendServices/456"
    key = ec.generate_private_key(ec.SECP256R1())
    jwk = json.loads(jwt.algorithms.ECAlgorithm.to_jwk(key.public_key()))
    jwk.update(kid="test", alg="ES256", use="sig")
    monkeypatch.setattr(id_token, "_fetch_certs", lambda *args: {"keys": [jwk]})
    signed = jwt.encode(
        {
            "iss": "https://cloud.google.com/iap",
            "aud": settings.iap_audience,
            "sub": "123",
            "email": "user@example.org",
            "iat": int(time.time()) - 5,
            "exp": int(time.time()) + 300,
        },
        key,
        algorithm="ES256",
        headers={"kid": "test"},
    )

    class SourceCollector(CountingCollector):
        def collect(self, pool):
            snapshot = super().collect(pool)
            snapshot["metrics"].update(
                source="GMP cAdvisor",
                memory_definition="Working set",
                cadence_seconds=60,
                cpu_rate_window_seconds=300,
                cpu_alignment_window_seconds=None,
                units={"cpu": "cores", "memory": "bytes"},
                query_definition={
                    "cpu_metric": "container_cpu_usage_seconds_total",
                    "memory_metric": "container_memory_working_set_bytes",
                    "cpu_aggregation": "sum containers after rate[5m]",
                    "memory_aggregation": "sum container samples",
                },
                resource_status={"cpu": "available", "memory": "unavailable"},
                fallback_reason="Native query failed; selected validated GMP partial observations",
            )
            for point in snapshot["metrics"]["samples"]:
                point.pop("memory", None)
            return snapshot

    collector = SourceCollector(settings)
    with TestClient(create_app(settings, collector), headers={"x-goog-iap-jwt-assertion": signed}) as client:
        assert client.get("/api/inventory", headers={"x-goog-iap-jwt-assertion": ""}).status_code == 401
        pool = initialized(client)["pools"][0]
        before = collector.calls
        for days in [1, 5, 10]:
            usage = client.get("/api/snapshot", params={"pool_id": pool["id"], "days": days}).json()["usage"]
            assert usage["cpu_rate_window_seconds"] == 300
            assert usage["resource_status"] == {"cpu": "available", "memory": "unavailable"}
            assert usage["query_definition"]["memory_metric"] == "container_memory_working_set_bytes"
        for table in ["workloads", "pods"]:
            for fmt in ["csv", "xlsx"]:
                response = client.get(f"/api/export/{table}", params={"pool_id": pool["id"], "format": fmt})
                if fmt == "csv":
                    row = next(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
                else:
                    values = list(load_workbook(io.BytesIO(response.content)).active.values)
                    row = dict(zip(values[0], values[1], strict=True))
                assert row["Metric source"] == "GMP cAdvisor"
                assert str(row["CPU rate window seconds"]) == "300"
                assert row["Memory observations"] == "unavailable"
                assert row["Fallback reason"] == usage["fallback_reason"]
        assert collector.calls == before
